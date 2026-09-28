from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from PySide6.QtCore import Qt, QTimer, Signal
from PySide6.QtGui import QFont, QPixmap, QWheelEvent, QTextCursor
from PySide6.QtWidgets import (
    QAbstractItemView,
    QFileDialog,
    QFrame,
    QHeaderView,
    QHBoxLayout,
    QLabel,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QSplitter,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QTreeWidget,
    QTreeWidgetItem,
    QVBoxLayout,
    QWidget,
)

from ui.widgets.detachable_panel import DetachablePanel


class ZoomableImageLabel(QLabel):
    """Image widget that supports smooth, predictable zoom steps."""

    zoom_requested = Signal(int)

    def __init__(self, parent=None):
        super().__init__(parent)

        self.setAlignment(Qt.AlignCenter)
        self.setMinimumSize(1, 1)
        self.setText("No image selected")
        self.setMouseTracking(True)

    def wheelEvent(self, event: QWheelEvent):
        # Ctrl + mouse wheel = zoom.
        # Normal mouse wheel remains available to QScrollArea for scrolling.
        if event.modifiers() & Qt.ControlModifier:
            delta = event.angleDelta().y()
            if delta > 0:
                self.zoom_requested.emit(1)
            elif delta < 0:
                self.zoom_requested.emit(-1)
            event.accept()
            return

        event.ignore()


class ImagePreviewWidget(QWidget):
    """Preview area with fit-to-window behavior and 10% zoom steps."""

    def __init__(self, parent=None):
        super().__init__(parent)

        self.original_pixmap = QPixmap()
        self.fit_scale = 1.0
        self.zoom_level = 1.0
        self.zoom_step = 0.10

        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        toolbar = QFrame()
        toolbar.setObjectName("ImagePreviewToolbar")

        toolbar_layout = QHBoxLayout(toolbar)
        toolbar_layout.setContentsMargins(8, 6, 8, 6)
        toolbar_layout.setSpacing(6)

        self.fit_button = QPushButton("Fit")
        self.zoom_out_button = QPushButton("−")
        self.zoom_in_button = QPushButton("+")

        self.zoom_out_button.setToolTip("Zoom out 10%")
        self.zoom_in_button.setToolTip("Zoom in 10%")
        self.fit_button.setToolTip("Fit the complete image to the viewer")

        self.zoom_label = QLabel("100%")
        self.zoom_label.setObjectName("ImageZoomLabel")
        self.zoom_label.setMinimumWidth(58)
        self.zoom_label.setAlignment(Qt.AlignCenter)

        toolbar_layout.addWidget(self.fit_button)
        toolbar_layout.addWidget(self.zoom_out_button)
        toolbar_layout.addWidget(self.zoom_label)
        toolbar_layout.addWidget(self.zoom_in_button)
        toolbar_layout.addStretch()

        self.image_scroll = QScrollArea()
        self.image_scroll.setWidgetResizable(True)
        self.image_scroll.setAlignment(Qt.AlignCenter)
        self.image_scroll.setFrameShape(QFrame.NoFrame)
        self.image_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.image_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)

        self.image_label = ZoomableImageLabel()
        self.image_scroll.setWidget(self.image_label)

        layout.addWidget(toolbar)
        layout.addWidget(self.image_scroll, stretch=1)

        self.fit_button.clicked.connect(self.fit_image)
        self.zoom_out_button.clicked.connect(lambda: self.change_zoom(-1))
        self.zoom_in_button.clicked.connect(lambda: self.change_zoom(1))
        self.image_label.zoom_requested.connect(self.change_zoom)

    def set_image(self, pixmap: QPixmap):
        self.original_pixmap = pixmap
        self.zoom_level = 1.0
        self.update_zoom_label()

        self.image_label.setPixmap(pixmap)
        self.image_label.adjustSize()

        QTimer.singleShot(0, self.fit_image)

    def clear_image(self):
        self.original_pixmap = QPixmap()
        self.fit_scale = 1.0
        self.zoom_level = 1.0
        self.image_label.clear()
        self.image_label.setText("No image selected")
        self.update_zoom_label()

    def fit_image(self):
        if self.original_pixmap.isNull():
            return

        viewport = self.image_scroll.viewport().size()
        viewport_width = max(1, viewport.width() - 6)
        viewport_height = max(1, viewport.height() - 6)

        image_width = max(1, self.original_pixmap.width())
        image_height = max(1, self.original_pixmap.height())

        width_scale = viewport_width / image_width
        height_scale = viewport_height / image_height

        self.fit_scale = max(0.05, min(width_scale, height_scale))
        self.zoom_level = 1.0

        self.apply_scale()

    def change_zoom(self, direction: int):
        if self.original_pixmap.isNull():
            return

        next_zoom = self.zoom_level + (
            self.zoom_step if direction > 0 else -self.zoom_step
        )

        # Keep zoom deliberate rather than jumping too far.
        self.zoom_level = max(0.20, min(5.00, next_zoom))
        self.apply_scale()

    def apply_scale(self):
        if self.original_pixmap.isNull():
            return

        effective_scale = self.fit_scale * self.zoom_level

        target_width = max(
            1,
            round(self.original_pixmap.width() * effective_scale),
        )
        target_height = max(
            1,
            round(self.original_pixmap.height() * effective_scale),
        )

        scaled = self.original_pixmap.scaled(
            target_width,
            target_height,
            Qt.KeepAspectRatio,
            Qt.SmoothTransformation,
        )

        self.image_label.setPixmap(scaled)
        self.image_label.adjustSize()
        self.update_zoom_label()

    def update_zoom_label(self):
        if self.original_pixmap.isNull():
            self.zoom_label.setText("100%")
            return

        percentage = round(self.zoom_level * 100)
        self.zoom_label.setText(f"{percentage}%")

    def resizeEvent(self, event):
        super().resizeEvent(event)

        if self.original_pixmap.isNull():
            return

        # Recompute the fit baseline whenever the docked/detached window
        # changes size. A 130% zoom, for example, remains 130% of the new fit.
        old_zoom = self.zoom_level
        viewport = self.image_scroll.viewport().size()
        viewport_width = max(1, viewport.width() - 6)
        viewport_height = max(1, viewport.height() - 6)

        image_width = max(1, self.original_pixmap.width())
        image_height = max(1, self.original_pixmap.height())

        self.fit_scale = max(
            0.05,
            min(
                viewport_width / image_width,
                viewport_height / image_height,
            ),
        )
        self.zoom_level = old_zoom
        self.apply_scale()



class ZoomableTextEdit(QTextEdit):
    """Read-only text editor with Ctrl+wheel zoom support."""

    zoom_requested = Signal(int)

    def wheelEvent(self, event: QWheelEvent):
        if event.modifiers() & Qt.ControlModifier:
            delta = event.angleDelta().y()
            if delta > 0:
                self.zoom_requested.emit(1)
            elif delta < 0:
                self.zoom_requested.emit(-1)
            event.accept()
            return

        super().wheelEvent(event)


class TextPreviewWidget(QWidget):
    """Text preview with the same deliberate 10% zoom behavior as images."""

    def __init__(self, parent=None):
        super().__init__(parent)

        self.zoom_level = 1.0
        self.zoom_step = 0.10
        self.base_font = QFont("Consolas", 10)

        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        toolbar = QFrame()
        toolbar.setObjectName("TextPreviewToolbar")

        toolbar_layout = QHBoxLayout(toolbar)
        toolbar_layout.setContentsMargins(8, 6, 8, 6)
        toolbar_layout.setSpacing(6)

        self.fit_button = QPushButton("Fit")
        self.zoom_out_button = QPushButton("−")
        self.zoom_in_button = QPushButton("+")

        self.fit_button.setToolTip("Reset text size to 100%")
        self.zoom_out_button.setToolTip("Reduce text size by 10%")
        self.zoom_in_button.setToolTip("Increase text size by 10%")

        self.zoom_label = QLabel("100%")
        self.zoom_label.setObjectName("TextZoomLabel")
        self.zoom_label.setMinimumWidth(58)
        self.zoom_label.setAlignment(Qt.AlignCenter)

        toolbar_layout.addWidget(self.fit_button)
        toolbar_layout.addWidget(self.zoom_out_button)
        toolbar_layout.addWidget(self.zoom_label)
        toolbar_layout.addWidget(self.zoom_in_button)
        toolbar_layout.addStretch()

        self.editor = ZoomableTextEdit()
        self.editor.setReadOnly(True)
        self.editor.setLineWrapMode(QTextEdit.NoWrap)
        self.editor.setFont(self.base_font)

        layout.addWidget(toolbar)
        layout.addWidget(self.editor, stretch=1)

        self.fit_button.clicked.connect(self.reset_zoom)
        self.zoom_out_button.clicked.connect(lambda: self.change_zoom(-1))
        self.zoom_in_button.clicked.connect(lambda: self.change_zoom(1))
        self.editor.zoom_requested.connect(self.change_zoom)

    def set_text(self, text: str):
        self.editor.setPlainText(text)
        self.editor.moveCursor(QTextCursor.Start)
        self.reset_zoom()

    def reset_zoom(self):
        self.zoom_level = 1.0
        self.apply_zoom()

    def change_zoom(self, direction: int):
        next_zoom = self.zoom_level + (
            self.zoom_step if direction > 0 else -self.zoom_step
        )
        self.zoom_level = max(0.50, min(3.00, next_zoom))
        self.apply_zoom()

    def apply_zoom(self):
        font = QFont(self.base_font)
        font.setPointSizeF(self.base_font.pointSizeF() * self.zoom_level)
        self.editor.setFont(font)
        self.zoom_label.setText(f"{round(self.zoom_level * 100)}%")


class DataExplorerPage(QWidget):
    """
    VS Code-style browser for the generated Processed_Data hierarchy.

    Features:
      - Recursive Processed_Data hierarchy
      - TXT/CSV/LOG/MD preview with zoom controls
      - Excel preview
      - Image preview with automatic fit-to-window
      - 10% zoom in/out controls
      - Ctrl + mouse-wheel zoom
      - Detachable file-preview panel
    """

    DATA_ROLE = Qt.UserRole + 1
    TYPE_ROLE = Qt.UserRole + 2

    def __init__(self):
        super().__init__()

        self.root_path = Path("Processed_Data").resolve()
        self.current_file: Path | None = None

        self.setup_ui()
        self.refresh_tree()

    # ==========================================================
    # UI
    # ==========================================================

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 16, 20, 12)
        layout.setSpacing(10)

        title = QLabel("Processed Data")
        title.setObjectName("PageTitle")

        subtitle = QLabel(
            "Browse the generated Processed_Data hierarchy and preview files "
            "directly inside the DRSP workspace."
        )
        subtitle.setObjectName("PageSubtitle")
        subtitle.setWordWrap(True)

        layout.addWidget(title)
        layout.addWidget(subtitle)

        # ------------------------------------------------------
        # TOP TOOLBAR
        # ------------------------------------------------------

        toolbar = QHBoxLayout()
        toolbar.setSpacing(8)

        self.root_label = QLabel(str(self.root_path))
        self.root_label.setObjectName("ExplorerPathLabel")
        self.root_label.setWordWrap(False)

        self.refresh_button = QPushButton("Refresh")
        self.refresh_button.clicked.connect(self.refresh_tree)

        self.choose_root_button = QPushButton("Choose Root")
        self.choose_root_button.clicked.connect(self.choose_root)

        self.expand_button = QPushButton("Expand All")
        self.expand_button.clicked.connect(self.tree_expand_all)

        self.collapse_button = QPushButton("Collapse All")
        self.collapse_button.clicked.connect(self.tree_collapse_all)

        toolbar.addWidget(QLabel("Root:"))
        toolbar.addWidget(self.root_label, stretch=1)
        toolbar.addWidget(self.choose_root_button)
        toolbar.addWidget(self.refresh_button)
        toolbar.addWidget(self.expand_button)
        toolbar.addWidget(self.collapse_button)

        layout.addLayout(toolbar)

        # ------------------------------------------------------
        # MAIN EXPLORER SPLITTER
        # ------------------------------------------------------

        splitter = QSplitter(Qt.Horizontal)
        splitter.setHandleWidth(8)
        splitter.setChildrenCollapsible(True)

        # ------------------------------------------------------
        # HIERARCHY
        # ------------------------------------------------------

        hierarchy_panel = QWidget()
        hierarchy_layout = QVBoxLayout(hierarchy_panel)
        hierarchy_layout.setContentsMargins(0, 0, 0, 0)
        hierarchy_layout.setSpacing(6)

        hierarchy_title = QLabel("EXPLORER")
        hierarchy_title.setObjectName("ExplorerSectionTitle")

        self.tree = QTreeWidget()
        self.tree.setHeaderHidden(True)
        self.tree.setIndentation(18)
        self.tree.setAnimated(True)
        self.tree.setUniformRowHeights(True)
        self.tree.setSelectionMode(QAbstractItemView.SingleSelection)
        self.tree.itemClicked.connect(self.handle_item_clicked)

        hierarchy_layout.addWidget(hierarchy_title)
        hierarchy_layout.addWidget(self.tree, stretch=1)

        # ------------------------------------------------------
        # VIEWER CONTENT
        # ------------------------------------------------------

        viewer_content = QWidget()
        viewer_layout = QVBoxLayout(viewer_content)
        viewer_layout.setContentsMargins(0, 0, 0, 0)
        viewer_layout.setSpacing(0)

        viewer_header = QWidget()
        viewer_header_layout = QVBoxLayout(viewer_header)
        viewer_header_layout.setContentsMargins(10, 6, 10, 6)
        viewer_header_layout.setSpacing(2)

        self.viewer_name_label = QLabel("No file selected")
        self.viewer_name_label.setObjectName("ExplorerViewerTitle")

        self.viewer_path_label = QLabel(
            "Select a file from the hierarchy."
        )
        self.viewer_path_label.setObjectName("ExplorerViewerPath")
        self.viewer_path_label.setWordWrap(True)

        viewer_header_layout.addWidget(self.viewer_name_label)
        viewer_header_layout.addWidget(self.viewer_path_label)

        self.viewer_stack = QStackedWidget()

        self.empty_view = self.create_empty_view()
        self.text_preview = TextPreviewWidget()
        self.text_view = self.text_preview
        self.image_preview = ImagePreviewWidget()
        self.image_view = self.image_preview
        self.workbook_view = self.create_workbook_view()
        self.unsupported_view = self.create_unsupported_view()

        self.viewer_stack.addWidget(self.empty_view)
        self.viewer_stack.addWidget(self.text_view)
        self.viewer_stack.addWidget(self.image_view)
        self.viewer_stack.addWidget(self.workbook_view)
        self.viewer_stack.addWidget(self.unsupported_view)

        viewer_layout.addWidget(viewer_header)
        viewer_layout.addWidget(self.viewer_stack, stretch=1)

        # ------------------------------------------------------
        # DETACHABLE FILE VIEWER
        # ------------------------------------------------------

        self.viewer_panel = DetachablePanel(
            "File Viewer",
            viewer_content,
        )
        self.viewer_panel.setMinimumSize(0, 0)

        splitter.addWidget(hierarchy_panel)
        splitter.addWidget(self.viewer_panel)

        splitter.setSizes([330, 830])
        splitter.setStretchFactor(0, 0)
        splitter.setStretchFactor(1, 1)

        layout.addWidget(splitter, stretch=1)

    # ==========================================================
    # EMPTY VIEW
    # ==========================================================

    def create_empty_view(self):
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setAlignment(Qt.AlignCenter)

        title = QLabel("Processed Data Explorer")
        title.setObjectName("ExplorerEmptyTitle")
        title.setAlignment(Qt.AlignCenter)

        text = QLabel(
            "Select a TXT, PNG, or Excel file from the hierarchy\n"
            "to preview it here."
        )
        text.setAlignment(Qt.AlignCenter)
        text.setObjectName("ExplorerEmptyText")

        layout.addWidget(title)
        layout.addWidget(text)
        return panel

    # ==========================================================
    # TEXT VIEW
    # ==========================================================

    def create_text_view(self):
        return TextPreviewWidget()

    # ==========================================================
    # WORKBOOK VIEW
    # ==========================================================

    def create_workbook_view(self):
        self.workbook_tabs = QStackedWidget()
        return self.workbook_tabs

    # ==========================================================
    # UNSUPPORTED VIEW
    # ==========================================================

    def create_unsupported_view(self):
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setAlignment(Qt.AlignCenter)

        self.unsupported_label = QLabel()
        self.unsupported_label.setAlignment(Qt.AlignCenter)
        self.unsupported_label.setWordWrap(True)

        layout.addWidget(self.unsupported_label)
        return panel

    # ==========================================================
    # ROOT SELECTION
    # ==========================================================

    def choose_root(self):
        selected = QFileDialog.getExistingDirectory(
            self,
            "Select Processed_Data Root Folder",
            str(self.root_path if self.root_path.exists() else Path.cwd()),
        )

        if not selected:
            return

        self.root_path = Path(selected).resolve()
        self.root_label.setText(str(self.root_path))
        self.refresh_tree()

    # ==========================================================
    # TREE REFRESH
    # ==========================================================

    def refresh_tree(self):
        self.tree.clear()

        if not self.root_path.exists():
            root_item = QTreeWidgetItem(["Processed_Data — not found"])
            root_item.setData(0, self.DATA_ROLE, self.root_path)
            root_item.setData(0, self.TYPE_ROLE, "missing")
            self.tree.addTopLevelItem(root_item)

            self.viewer_name_label.setText("Processed_Data not found")
            self.viewer_path_label.setText(
                f"Expected location: {self.root_path}"
            )
            self.viewer_stack.setCurrentWidget(self.empty_view)
            self.image_preview.clear_image()
            return

        root_item = QTreeWidgetItem([
            self.root_path.name or "Processed_Data"
        ])
        root_item.setData(0, self.DATA_ROLE, self.root_path)
        root_item.setData(0, self.TYPE_ROLE, "directory")
        self.tree.addTopLevelItem(root_item)

        self.add_directory_children(root_item, self.root_path)
        root_item.setExpanded(True)

        self.viewer_stack.setCurrentWidget(self.empty_view)
        self.viewer_name_label.setText("No file selected")
        self.viewer_path_label.setText(
            "Select a file from the hierarchy."
        )
        self.image_preview.clear_image()

    def add_directory_children(
        self,
        parent_item: QTreeWidgetItem,
        directory: Path,
    ):
        try:
            entries = sorted(
                directory.iterdir(),
                key=lambda item: (
                    not item.is_dir(),
                    item.name.lower(),
                ),
            )
        except OSError:
            return

        for entry in entries:
            if entry.name.startswith("."):
                continue

            item_type = "directory" if entry.is_dir() else "file"
            child = QTreeWidgetItem([entry.name])
            child.setData(0, self.DATA_ROLE, entry)
            child.setData(0, self.TYPE_ROLE, item_type)

            if entry.is_dir():
                child.setChildIndicatorPolicy(
                    QTreeWidgetItem.ShowIndicator
                )
                self.add_directory_children(child, entry)

            parent_item.addChild(child)

    # ==========================================================
    # TREE ACTIONS
    # ==========================================================

    def tree_expand_all(self):
        self.tree.expandAll()

    def tree_collapse_all(self):
        self.tree.collapseAll()

        root = self.tree.topLevelItem(0)
        if root is not None:
            root.setExpanded(True)

    def handle_item_clicked(
        self,
        item: QTreeWidgetItem,
        _column: int,
    ):
        path = item.data(0, self.DATA_ROLE)
        item_type = item.data(0, self.TYPE_ROLE)

        if not path or item_type == "directory":
            return

        self.open_file(Path(path))

    # ==========================================================
    # FILE OPENING
    # ==========================================================

    def open_file(self, path: Path):
        self.current_file = path

        self.viewer_name_label.setText(path.name)
        try:
            relative = path.relative_to(self.root_path)
            self.viewer_path_label.setText(str(relative))
        except ValueError:
            self.viewer_path_label.setText(str(path))

        suffix = path.suffix.lower()

        try:
            if suffix in {".txt", ".csv", ".log", ".md"}:
                self.show_text_file(path)
                self.log_opened_file(path)
                return

            if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".webp"}:
                self.show_image_file(path)
                self.log_opened_file(path)
                return

            if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
                self.show_workbook(path)
                self.log_opened_file(path)
                return

            self.unsupported_label.setText(
                f"No embedded preview is available for:\n\n{path.name}\n\n"
                "The file remains available in the Processed_Data hierarchy."
            )
            self.viewer_stack.setCurrentWidget(self.unsupported_view)
            self.log_opened_file(path)

        except Exception as exc:
            QMessageBox.critical(
                self,
                "Preview Error",
                f"Could not preview:\n{path}\n\n{exc}",
            )

    def show_text_file(self, path: Path):
        content = path.read_text(
            encoding="utf-8",
            errors="replace",
        )

        self.text_preview.set_text(content)
        self.viewer_stack.setCurrentWidget(self.text_view)

    def show_image_file(self, path: Path):
        pixmap = QPixmap(str(path))

        if pixmap.isNull():
            raise RuntimeError("The selected image could not be decoded.")

        self.image_preview.set_image(pixmap)
        self.viewer_stack.setCurrentWidget(self.image_view)

    def show_workbook(self, path: Path):
        while self.workbook_tabs.count():
            widget = self.workbook_tabs.widget(0)
            self.workbook_tabs.removeWidget(widget)
            widget.deleteLater()

        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
        )

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]

            table = QTableWidget()
            table.setEditTriggers(QAbstractItemView.NoEditTriggers)
            table.setSelectionBehavior(QAbstractItemView.SelectRows)
            table.setSelectionMode(QAbstractItemView.SingleSelection)
            table.setAlternatingRowColors(True)
            table.setWordWrap(False)
            table.setColumnCount(0)
            table.setRowCount(0)

            rows = list(
                worksheet.iter_rows(
                    values_only=True,
                )
            )

            if not rows:
                self.workbook_tabs.addWidget(table)
                continue

            column_count = max(len(row) for row in rows)

            table.setColumnCount(column_count)
            table.setRowCount(len(rows))

            header_values = rows[0]
            headers = []
            for index in range(column_count):
                value = (
                    header_values[index]
                    if index < len(header_values)
                    else ""
                )
                headers.append(str(value) if value is not None else "")

            table.setHorizontalHeaderLabels(headers)

            for row_index, row in enumerate(rows):
                for column_index in range(column_count):
                    value = (
                        row[column_index]
                        if column_index < len(row)
                        else ""
                    )
                    item = QTableWidgetItem(
                        "" if value is None else str(value)
                    )
                    item.setTextAlignment(Qt.AlignCenter)
                    table.setItem(
                        row_index,
                        column_index,
                        item,
                    )

            table.horizontalHeader().setSectionResizeMode(
                QHeaderView.Interactive
            )

            self.workbook_tabs.addWidget(table)

        workbook.close()
        self.viewer_stack.setCurrentWidget(self.workbook_view)

    # ==========================================================
    # CONSOLE INTEGRATION
    # ==========================================================

    def log_opened_file(self, path: Path):
        window = self.window()
        if hasattr(window, "console_text"):
            window.console_text.append(
                f"[DRSP] Opened processed data: {path}"
            )
