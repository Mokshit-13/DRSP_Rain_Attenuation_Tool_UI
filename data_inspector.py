"""
================================================================================
DRSP Rain Attenuation Tool — Data Inspector
================================================================================
Standalone reverse-engineering utility for unknown / unfamiliar raw NARL
datasets.

Purpose
-------
Before writing or adjusting any processing algorithm, this tool inspects a
raw dataset (single file, single month, or entire year) and reports its
structure, timing characteristics, channel statistics, data quality, and
likely compatibility with the DRSP Rain Attenuation Tool.

This utility NEVER modifies, repairs, or removes anything.  It only reads
and reports observations.
================================================================================
"""

from __future__ import annotations

import re
import sys
import tkinter as tk
from collections import Counter
from datetime import datetime
from pathlib import Path
from tkinter import filedialog

import pandas as pd


# ==============================================================================
# CONFIGURATION
# ==============================================================================

OUTPUT_ROOT = "Processed_Data"
REPORT_DIR  = "Inspection_Reports"

REFERENCE_COLUMN_HINTS = ["ref", "reference", "baseline"]


# ==============================================================================
# TERMINAL HELPERS
# ==============================================================================

_SEP = "=" * 57


def _header(title: str) -> None:
    print()
    print(_SEP)
    print(title.center(57))
    print(_SEP)


# ==============================================================================
# UI — MENUS & FOLDER/FILE SELECTION
# ==============================================================================

def show_menu() -> str:
    """Displays the main menu and returns the user's validated choice."""
    _header("DATA INSPECTOR")
    print()
    print("  Select Inspection Mode")
    print()
    print("    1.  Inspect Single File")
    print("    2.  Inspect Single Month")
    print("    3.  Inspect Entire Year")
    print()
    print("    0.  Exit")
    print()

    while True:
        choice = input("  Enter Choice : ").strip()
        if choice in ("0", "1", "2", "3"):
            return choice
        print("  Invalid choice. Please enter 0, 1, 2, or 3.")


def select_report_format() -> str:
    """Displays the report format menu and returns '1' (TXT) or '2' (Excel)."""
    print()
    print(_SEP)
    print()
    print("  Select Report Format")
    print()
    print("    1.  TXT Report")
    print("    2.  Excel Report (.xlsx)")
    print()

    while True:
        choice = input("  Enter Choice : ").strip()
        if choice in ("1", "2"):
            return choice
        print("  Invalid choice. Please enter 1 or 2.")


def _open_folder_dialog(title: str) -> str:
    """Opens a standard Windows folder-selection dialog. Exits on cancel."""
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askdirectory(title=title)
    root.destroy()

    if not path:
        print("No folder selected. Exiting.")
        sys.exit(0)

    return path


def _open_file_dialog(title: str) -> str:
    """Opens a standard Windows file-selection dialog. Exits on cancel."""
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title=title,
        filetypes=[("Text files", "*.txt")],
    )
    root.destroy()

    if not path:
        print("No file selected. Exiting.")
        sys.exit(0)

    return path


def _find_raw_files(folder: Path) -> list[Path]:
    """
    Finds every candidate raw NARL/NAR data file inside a folder
    (non-recursive) — any *.txt file.
    """
    return sorted(folder.glob("*.txt"))


# ==============================================================================
# FILE-LEVEL INSPECTION
# ==============================================================================

def _guess_delimiter(raw_lines: list[str]) -> str:
    """
    Examines the first non-empty line and guesses the delimiter.
    Returns one of: 'whitespace', 'tab', 'comma', 'unknown'.
    """
    for line in raw_lines:
        if line.strip():
            if "\t" in line:
                return "tab"
            if "," in line:
                return "comma"
            if re.search(r"\s{2,}|\s", line.strip()):
                return "whitespace"
            return "unknown"
    return "unknown"


def _guess_timestamp_format(sample_values: list[str]) -> str:
    """
    Examines a sample of raw Time-column string values and makes an
    educated guess at the timestamp format.
    """
    if not sample_values:
        return "Unrecognised (no data)"

    patterns = [
        (r"^\d{2}:\d{2}:\d{2}\.\d+$", "HH:MM:SS.sss (colon-separated with fractional seconds)"),
        (r"^\d{2}:\d{2}:\d{2}$",       "HH:MM:SS (colon-separated)"),
        (r"^\d{8}$",                   "HHMMSSCC (8-digit, possible centiseconds)"),
        (r"^\d{6}$",                   "HHMMSS (6-digit, no separators)"),
    ]

    counts = Counter()
    for pattern, label in patterns:
        matches = sum(1 for v in sample_values if re.match(pattern, str(v).strip()))
        if matches:
            counts[label] = matches

    if not counts:
        return f"Unrecognised — example value: '{sample_values[0]}'"

    best_label, _ = counts.most_common(1)[0]
    return best_label


def _classify_filename(file_path: Path) -> str:
    """
    Attempts to classify the dataset generation/source based on the
    filename pattern.
    """
    name = file_path.stem

    if re.match(r"^NARL_\d{1,2}_\d{1,2}_\d{4}$", name):
        return "NARL-series (post-2020 style naming: NARL_D_M_YYYY)"
    if re.match(r"^NAR_\d{1,2}_\d{1,2}_\d{4}$", name):
        return "NAR-series (legacy style naming: NAR_D_M_YYYY)"
    if re.match(r"^\d{8}$", name):
        return "Numeric date stem (possible DDMMYYYY or YYYYMMDD)"

    return f"Unrecognised naming pattern — example: '{file_path.name}'"


def _identify_channels(columns: list[str]) -> tuple[list[str], list[str]]:
    """
    Separates DataFrame columns into likely data channels and likely
    reference channels based on common naming hints.  Any column matching
    neither hint set (and not 'Time') is treated as a data channel by
    default, since most NARL datasets only carry amplitude columns.
    """
    data_channels = []
    reference_channels = []

    for col in columns:
        lowered = col.lower()
        if lowered == "time":
            continue
        if any(hint in lowered for hint in REFERENCE_COLUMN_HINTS):
            reference_channels.append(col)
        else:
            data_channels.append(col)

    return data_channels, reference_channels


def inspect_file(file_path: Path) -> dict:
    """
    Performs a full inspection of a single raw data file.

    Returns a dict containing every category of observation:
        general, time, channels, quality, structure, filename_classification,
        compatibility
    Never raises on malformed files — captures an "error" key instead.
    """
    file_path = Path(file_path)
    report = {
        "file_path": str(file_path),
        "error": None,
    }

    try:
        file_size = file_path.stat().st_size
    except OSError:
        file_size = None

    try:
        raw_lines = file_path.read_text(errors="replace").splitlines()
    except Exception as exc:
        report["error"] = f"Could not read file: {exc}"
        return report

    delimiter = _guess_delimiter(raw_lines)
    header_present = bool(raw_lines) and not re.match(r"^[\d\.\-\+:\s]+$", raw_lines[0].strip())

    try:
        df = pd.read_csv(file_path, sep=r"\s+", engine="python")
        read_ok = True
    except Exception as exc:
        df = None
        read_ok = False
        report["error"] = f"Could not parse as tabular data: {exc}"

    num_rows = len(df) if read_ok else 0
    num_cols = len(df.columns) if read_ok else 0
    columns  = list(df.columns) if read_ok else []

    # ── GENERAL ─────────────────────────────────────────────────────────────
    report["general"] = {
        "File Name"         : file_path.name,
        "File Size (bytes)" : file_size,
        "Number of Rows"    : num_rows,
        "Number of Columns" : num_cols,
        "Header Present"    : "Yes" if header_present else "No",
        "Delimiter"         : delimiter,
    }

    # ── TIME ANALYSIS ────────────────────────────────────────────────────────
    time_info = {}
    if read_ok and "Time" in df.columns:
        raw_time_str = df["Time"].astype(str).str.strip()
        sample_vals   = raw_time_str.head(20).tolist()

        parsed = pd.to_datetime(raw_time_str, errors="coerce", format="%H:%M:%S")
        valid_mask = parsed.notna()

        n_valid    = int(valid_mask.sum())
        n_invalid  = int((~valid_mask).sum())
        n_unique   = int(raw_time_str.nunique())
        n_dupes    = int(len(raw_time_str) - n_unique)

        strictly_increasing = None
        min_step = max_step = sampling_interval = None

        if n_valid > 1:
            valid_times = parsed[valid_mask].reset_index(drop=True)
            diffs = valid_times.diff().dropna().dt.total_seconds()
            if not diffs.empty:
                min_step = float(diffs.min())
                max_step = float(diffs.max())
                mode_vals = diffs.mode()
                sampling_interval = float(mode_vals.iloc[0]) if not mode_vals.empty else None
                strictly_increasing = bool((diffs > 0).all())

        time_info = {
            "First Timestamp"           : str(raw_time_str.iloc[0]) if len(raw_time_str) else "N/A",
            "Last Timestamp"            : str(raw_time_str.iloc[-1]) if len(raw_time_str) else "N/A",
            "Unique Timestamp Count"    : n_unique,
            "Duplicate Timestamp Count" : n_dupes,
            "Invalid Timestamp Count"   : n_invalid,
            "Sampling Interval (s)"     : sampling_interval,
            "Minimum Time Step (s)"     : min_step,
            "Maximum Time Step (s)"     : max_step,
            "Strictly Increasing"       : (
                "Yes" if strictly_increasing else "No" if strictly_increasing is not None else "N/A"
            ),
            "Guessed Format"            : _guess_timestamp_format(sample_vals),
        }
    else:
        time_info = {"Note": "No 'Time' column found — time analysis skipped."}

    report["time"] = time_info

    # ── CHANNEL ANALYSIS ─────────────────────────────────────────────────────
    channel_stats = {}
    data_channels, reference_channels = ([], [])

    if read_ok:
        data_channels, reference_channels = _identify_channels(columns)

        for ch in data_channels + reference_channels:
            numeric = pd.to_numeric(df[ch], errors="coerce")
            finite  = numeric.replace([float("inf"), float("-inf")], pd.NA).dropna()

            if len(finite) > 0:
                channel_stats[ch] = {
                    "Minimum"            : float(finite.min()),
                    "Maximum"            : float(finite.max()),
                    "Mean"               : float(finite.mean()),
                    "Median"             : float(finite.median()),
                    "Standard Deviation" : float(finite.std()) if len(finite) > 1 else 0.0,
                    "Variance"           : float(finite.var()) if len(finite) > 1 else 0.0,
                }
            else:
                channel_stats[ch] = {
                    "Minimum": None, "Maximum": None, "Mean": None,
                    "Median": None, "Standard Deviation": None, "Variance": None,
                }

    report["channels"] = {
        "Number of Data Channels"      : len(data_channels),
        "Number of Reference Channels" : len(reference_channels),
        "Data Channel Names"           : data_channels,
        "Reference Channel Names"      : reference_channels,
        "Per-Channel Statistics"       : channel_stats,
    }

    # ── DATA QUALITY ─────────────────────────────────────────────────────────
    quality = {
        "INF Count"              : 0,
        "NaN Count"              : 0,
        "Blank Cell Count"       : 0,
        "Corrupted Rows"         : 0,
        "Rows With Wrong Column Count": 0,
        "Negative Values"        : 0,
        "Zero Values"            : 0,
        "Repeated Value Count"   : 0,
    }

    if read_ok:
        numeric_cols = data_channels + reference_channels
        for ch in numeric_cols:
            numeric = pd.to_numeric(df[ch], errors="coerce")
            quality["INF Count"] += int(numeric.isin([float("inf"), float("-inf")]).sum())
            quality["NaN Count"] += int(numeric.isna().sum())
            quality["Negative Values"] += int((numeric < 0).sum())
            quality["Zero Values"] += int((numeric == 0).sum())

            finite = numeric.replace([float("inf"), float("-inf")], pd.NA).dropna()
            if len(finite) > 0:
                value_counts = finite.value_counts()
                quality["Repeated Value Count"] += int((value_counts[value_counts > 1]).sum())

        blank_cells = int(df.isna().sum().sum()) - quality["NaN Count"]
        quality["Blank Cell Count"] = max(blank_cells, 0)

        expected_cols = num_cols
        wrong_count = 0
        lines_to_check = raw_lines[1:] if header_present else raw_lines
        for line in lines_to_check:
            if not line.strip():
                continue
            token_count = len(re.split(r"\s+", line.strip()))
            if token_count != expected_cols:
                wrong_count += 1
        quality["Rows With Wrong Column Count"] = wrong_count
        quality["Corrupted Rows"] = wrong_count

    report["quality"] = quality

    # ── STRUCTURE ────────────────────────────────────────────────────────────
    identical_structure = True
    if raw_lines:
        token_counts = {
            len(re.split(r"\s+", line.strip()))
            for line in raw_lines if line.strip()
        }
        identical_structure = len(token_counts) <= 1

    report["structure"] = {
        "Headers Exist"          : "Yes" if header_present else "No",
        "Appears Fixed-Width"    : "No" if delimiter in ("tab", "comma") else "Possibly",
        "Tab Separated"          : "Yes" if delimiter == "tab" else "No",
        "Comma Separated"        : "Yes" if delimiter == "comma" else "No",
        "Every Row Identical Structure": "Yes" if identical_structure else "No",
    }

    report["filename_classification"] = _classify_filename(file_path)

    n_data_ch = len(data_channels)
    n_ref_ch  = len(reference_channels)

    report["compatibility"] = {
        "Receiver Channels Detected" : n_data_ch,
        "Reference Channel Present"  : "Yes" if n_ref_ch > 0 else "No",
        "Single-Channel Compatible"  : "Yes" if n_data_ch >= 1 else "No",
        "Multi-Channel Compatible"   : "Yes" if n_data_ch >= 2 else "No",
    }

    return report


# ==============================================================================
# AGGREGATE INSPECTION (MONTH / YEAR)
# ==============================================================================

def inspect_month(month_folder) -> list:
    """
    Inspects every raw data file found inside every daily folder of a
    month.  Returns a list of per-file inspection dicts.
    """
    month_path    = Path(month_folder)
    daily_folders = sorted(f for f in month_path.iterdir() if f.is_dir())

    file_reports = []
    for day_folder in daily_folders:
        raw_files = _find_raw_files(day_folder)
        for raw_file in raw_files:
            report = inspect_file(raw_file)
            report["day_folder"] = day_folder.name
            file_reports.append(report)

    return file_reports


def inspect_year(year_folder) -> list:
    """
    Inspects every raw data file found inside every daily folder of every
    month in a year.  Returns a list of per-file inspection dicts, each
    tagged with its month and day folder.
    """
    year_path     = Path(year_folder)
    month_folders = sorted(f for f in year_path.iterdir() if f.is_dir())

    file_reports = []
    for month_folder in month_folders:
        month_reports = inspect_month(month_folder)
        for r in month_reports:
            r["month"] = month_folder.name
        file_reports.extend(month_reports)

    return file_reports


# ==============================================================================
# INSPECTION ENTRY POINTS
# ==============================================================================

def inspect_single_file() -> dict:
    """Inspects one raw data file selected by the user."""
    file_path = _open_file_dialog("Select Raw Data File to Inspect")
    report = inspect_file(Path(file_path))

    return {
        "mode"    : "file",
        "source"  : Path(file_path).parent.name,
        "reports" : [report],
    }


def inspect_single_month() -> dict:
    """Inspects every raw file inside every daily folder of one month."""
    month_folder = _open_folder_dialog("Select Month Folder to Inspect")
    month_name   = Path(month_folder).name
    reports      = inspect_month(month_folder)

    return {
        "mode"   : "month",
        "source" : month_name,
        "reports": reports,
    }


def inspect_entire_year() -> dict:
    """Inspects every raw file inside every month of one year."""
    year_folder = _open_folder_dialog("Select Year Folder to Inspect")
    year_name   = Path(year_folder).name
    reports     = inspect_year(year_folder)

    return {
        "mode"   : "year",
        "source" : year_name,
        "reports": reports,
    }


# ==============================================================================
# OBSERVATIONS
# ==============================================================================

def generate_observations(data: dict) -> list:
    """
    Builds a list of plain-language observations (not assumptions) based on
    the aggregate inspection results.  Only reports patterns actually seen
    in the data.
    """
    reports = [r for r in data["reports"] if not r.get("error")]
    observations = []

    if not reports:
        observations.append("No readable files were found — no observations available.")
        return observations

    formats_seen = Counter(
        r["time"].get("Guessed Format", "N/A")
        for r in reports if "time" in r and "Guessed Format" in r["time"]
    )
    if formats_seen:
        most_common_fmt, count = formats_seen.most_common(1)[0]
        observations.append(
            f"Most files ({count}/{len(reports)}) appear to use timestamp format: {most_common_fmt}"
        )
        if len(formats_seen) > 1:
            observations.append(
                f"Timestamp format is NOT fully consistent across files — {len(formats_seen)} distinct formats observed."
            )

    intervals = [
        r["time"].get("Sampling Interval (s)")
        for r in reports
        if r.get("time", {}).get("Sampling Interval (s)") is not None
    ]
    if intervals:
        interval_counts = Counter(intervals)
        common_interval, cnt = interval_counts.most_common(1)[0]
        observations.append(
            f"Most common sampling interval observed: {common_interval:g} second(s) "
            f"({cnt}/{len(intervals)} files)."
        )
        if len(interval_counts) > 1:
            observations.append(
                "Multiple distinct sampling intervals were found — dataset may require "
                "normalisation to a single sample rate before processing."
            )

    files_with_dupes = [
        r for r in reports
        if r.get("time", {}).get("Duplicate Timestamp Count", 0) > 0
    ]
    if files_with_dupes:
        observations.append(
            f"{len(files_with_dupes)}/{len(reports)} file(s) contain duplicate timestamps "
            "within the same second — may require averaging before analysis."
        )

    channel_counts = Counter(
        r["channels"]["Number of Data Channels"]
        for r in reports if "channels" in r
    )
    if channel_counts:
        common_count, cnt = channel_counts.most_common(1)[0]
        observations.append(
            f"Most files ({cnt}/{len(reports)}) contain {common_count} data channel(s)."
        )
        if len(channel_counts) > 1:
            observations.append(
                "Channel count is NOT consistent across all files — some files may use a "
                "different receiver configuration."
            )

    has_reference = sum(
        1 for r in reports
        if r.get("channels", {}).get("Number of Reference Channels", 0) > 0
    )
    if has_reference == 0:
        observations.append(
            "No dedicated reference channel column was detected in any file — reference "
            "levels are likely computed from the data channels themselves (e.g. top-N% method)."
        )
    else:
        observations.append(
            f"{has_reference}/{len(reports)} file(s) contain an explicit reference channel column."
        )

    total_inf = sum(r.get("quality", {}).get("INF Count", 0) for r in reports)
    total_nan = sum(r.get("quality", {}).get("NaN Count", 0) for r in reports)
    if total_inf > 0:
        observations.append(f"INF values detected across dataset: {total_inf} total occurrence(s).")
    if total_nan > 0:
        observations.append(f"NaN values detected across dataset: {total_nan} total occurrence(s).")

    wrong_col_files = [
        r for r in reports
        if r.get("quality", {}).get("Rows With Wrong Column Count", 0) > 0
    ]
    if wrong_col_files:
        observations.append(
            f"{len(wrong_col_files)}/{len(reports)} file(s) contain rows with an unexpected "
            "column count — possible corruption or malformed rows."
        )

    naming_patterns = Counter(r.get("filename_classification", "Unknown") for r in reports)
    if naming_patterns:
        common_pattern, cnt = naming_patterns.most_common(1)[0]
        observations.append(f"Dominant filename convention: {common_pattern} ({cnt}/{len(reports)} files).")

    if not observations:
        observations.append("No notable anomalies were observed during inspection.")

    return observations


# ==============================================================================
# REPORT HELPERS
# ==============================================================================

def _report_path(stem: str, extension: str) -> Path:
    """Builds and creates the Inspection_Reports output directory."""
    out_dir = Path(OUTPUT_ROOT) / REPORT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir / f"Inspection_Report_{stem}.{extension}"


def build_report_stem(data: dict) -> str:
    """
    Returns the filename stem (without extension) for any inspection report.

    Single File  → day-folder name with spaces replaced by underscores
    Single Month → month folder name as-is
    Entire Year  → year folder name as-is
    """
    source = data["source"]
    if data["mode"] == "file":
        return source.replace(" ", "_")
    return source


def _mode_label(data: dict) -> str:
    return {"file": "Single File", "month": "Single Month", "year": "Entire Year"}.get(
        data["mode"], data["mode"]
    )


def _aggregate_general_counts(data: dict) -> dict:
    """Computes dataset-wide totals for the General Information section."""
    reports = data["reports"]
    total_files = len(reports)
    total_rows  = sum(r.get("general", {}).get("Number of Rows", 0) for r in reports if not r.get("error"))
    total_size  = sum(
        r.get("general", {}).get("File Size (bytes)") or 0 for r in reports if not r.get("error")
    )
    errored     = sum(1 for r in reports if r.get("error"))

    return {
        "Total Files Scanned" : total_files,
        "Files With Errors"   : errored,
        "Total Rows (all files)": total_rows,
        "Total Size (bytes)"  : total_size,
    }


# ==============================================================================
# TXT REPORT
# ==============================================================================

def save_txt_report(data: dict) -> Path:
    """
    Generates a professionally formatted plain-text inspection report with
    clear section headers and simple tables.
    """
    reports   = data["reports"]
    source    = data["source"]
    generated = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    report_path = _report_path(build_report_stem(data), "txt")
    width = 78

    with open(report_path, "w", newline="", encoding="utf-8") as f:

        def w(line: str = "") -> None:
            f.write(line + "\n")

        def section(title: str) -> None:
            w()
            w("=" * width)
            w(title)
            w("=" * width)

        w("=" * width)
        w("DATA INSPECTION REPORT".center(width))
        w("=" * width)
        w()
        w(f"Inspection Mode  : {_mode_label(data)}")
        w(f"Dataset Selected : {source}")
        w(f"Generated On     : {generated}")

        section("GENERAL INFORMATION")
        agg = _aggregate_general_counts(data)
        for key, val in agg.items():
            w(f"  {key:<28}: {val}")

        w()
        w("  Per-File Summary")
        w("  " + "-" * (width - 2))
        for r in reports:
            name = r.get("general", {}).get("File Name", Path(r.get("file_path", "?")).name)
            if r.get("error"):
                w(f"    {name:<30} ERROR: {r['error']}")
                continue
            gen = r["general"]
            w(
                f"    {name:<30} rows={gen['Number of Rows']:<8} "
                f"cols={gen['Number of Columns']:<4} "
                f"delim={gen['Delimiter']:<10} header={gen['Header Present']}"
            )

        section("TIME ANALYSIS")
        for r in reports:
            if r.get("error"):
                continue
            name = r["general"]["File Name"]
            w(f"  File: {name}")
            for key, val in r.get("time", {}).items():
                w(f"    {key:<28}: {val}")
            w()

        section("CHANNEL ANALYSIS")
        for r in reports:
            if r.get("error"):
                continue
            name = r["general"]["File Name"]
            ch_info = r.get("channels", {})
            w(f"  File: {name}")
            w(f"    Data Channels      : {ch_info.get('Number of Data Channels', 0)} "
              f"({', '.join(ch_info.get('Data Channel Names', [])) or 'none'})")
            w(f"    Reference Channels : {ch_info.get('Number of Reference Channels', 0)} "
              f"({', '.join(ch_info.get('Reference Channel Names', [])) or 'none'})")

            stats = ch_info.get("Per-Channel Statistics", {})
            if stats:
                w("    Per-Channel Statistics")
                col_w = 14
                header = "    " + "Channel".ljust(20) + "".join(
                    s.ljust(col_w) for s in ["Min", "Max", "Mean", "Median", "StdDev", "Variance"]
                )
                w(header)
                for ch, s in stats.items():
                    row = "    " + ch.ljust(20)
                    for key in ("Minimum", "Maximum", "Mean", "Median", "Standard Deviation", "Variance"):
                        v = s.get(key)
                        row += (f"{v:.3f}".ljust(col_w) if isinstance(v, (int, float)) else "N/A".ljust(col_w))
                    w(row)
            w()

        section("DATA QUALITY")
        for r in reports:
            if r.get("error"):
                continue
            name = r["general"]["File Name"]
            w(f"  File: {name}")
            for key, val in r.get("quality", {}).items():
                w(f"    {key:<30}: {val}")
            w()

        section("DATASET STRUCTURE")
        for r in reports:
            if r.get("error"):
                continue
            name = r["general"]["File Name"]
            w(f"  File: {name}")
            for key, val in r.get("structure", {}).items():
                w(f"    {key:<32}: {val}")
            w(f"    {'Filename Classification':<32}: {r.get('filename_classification', 'N/A')}")
            w()

        section("COMPATIBILITY ANALYSIS")
        for r in reports:
            if r.get("error"):
                continue
            name = r["general"]["File Name"]
            w(f"  File: {name}")
            for key, val in r.get("compatibility", {}).items():
                w(f"    {key:<30}: {val}")
            w()

        section("INSPECTOR'S OBSERVATIONS")
        w()
        for obs in generate_observations(data):
            w(f"  • {obs}")

        w()
        w("=" * width)

    return report_path


# ==============================================================================
# EXCEL REPORT
# ==============================================================================

def save_excel_report(data: dict) -> Path:
    """
    Generates a professional multi-sheet Excel workbook:
        General Information | Time Analysis | Channel Statistics |
        Data Quality | Observations
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    reports = data["reports"]
    report_path = _report_path(build_report_stem(data), "xlsx")

    _thin = Side(style="thin")
    _border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
    _hdr_fill = PatternFill("solid", start_color="BDD7EE")

    def _hdr_style(cell) -> None:
        cell.font      = Font(name="Arial", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill      = _hdr_fill
        cell.border    = _border

    def _data_style(cell, align: str = "center") -> None:
        cell.font      = Font(name="Arial", size=10)
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.border    = _border

    def _write_table(ws, headers: list, rows: list, start_row: int = 1) -> int:
        """Writes a header row + data rows starting at start_row. Returns next free row."""
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=h)
            _hdr_style(cell)

        for r_idx, row_data in enumerate(rows, start=start_row + 1):
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=col_idx, value=val)
                _data_style(cell, align="left" if col_idx == 1 else "center")

        for col_idx, h in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            widths = [len(str(h))] + [
                len(str(row[col_idx - 1])) for row in rows if col_idx - 1 < len(row)
            ]
            ws.column_dimensions[col_letter].width = max(widths, default=10) + 4

        return start_row + len(rows) + 1

    wb = Workbook()

    # ── Sheet: General Information ────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "General Information"

    agg = _aggregate_general_counts(data)
    general_rows = [[k, v] for k, v in agg.items()]
    general_rows.append(["Inspection Mode", _mode_label(data)])
    general_rows.append(["Dataset Selected", data["source"]])
    general_rows.append(["Generated On", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])

    next_row = _write_table(ws1, ["Field", "Value"], general_rows, start_row=1)
    ws1.freeze_panes = "A2"

    ws1.cell(row=next_row + 1, column=1, value="Per-File Summary")
    per_file_headers = ["File Name", "Rows", "Columns", "Delimiter", "Header Present", "Error"]
    per_file_rows = []
    for r in reports:
        name = r.get("general", {}).get("File Name", Path(r.get("file_path", "?")).name)
        if r.get("error"):
            per_file_rows.append([name, "-", "-", "-", "-", r["error"]])
        else:
            gen = r["general"]
            per_file_rows.append([
                name, gen["Number of Rows"], gen["Number of Columns"],
                gen["Delimiter"], gen["Header Present"], "",
            ])
    _write_table(ws1, per_file_headers, per_file_rows, start_row=next_row + 2)

    # ── Sheet: Time Analysis ──────────────────────────────────────────────────
    ws2 = wb.create_sheet("Time Analysis")
    time_headers = [
        "File Name", "First Timestamp", "Last Timestamp", "Unique Count",
        "Duplicate Count", "Invalid Count", "Sampling Interval (s)",
        "Min Step (s)", "Max Step (s)", "Strictly Increasing", "Guessed Format",
    ]
    time_rows = []
    for r in reports:
        if r.get("error"):
            continue
        name = r["general"]["File Name"]
        t = r.get("time", {})
        time_rows.append([
            name,
            t.get("First Timestamp", "N/A"),
            t.get("Last Timestamp", "N/A"),
            t.get("Unique Timestamp Count", "N/A"),
            t.get("Duplicate Timestamp Count", "N/A"),
            t.get("Invalid Timestamp Count", "N/A"),
            t.get("Sampling Interval (s)", "N/A"),
            t.get("Minimum Time Step (s)", "N/A"),
            t.get("Maximum Time Step (s)", "N/A"),
            t.get("Strictly Increasing", "N/A"),
            t.get("Guessed Format", "N/A"),
        ])
    _write_table(ws2, time_headers, time_rows)
    ws2.freeze_panes = "A2"

    # ── Sheet: Channel Statistics ─────────────────────────────────────────────
    ws3 = wb.create_sheet("Channel Statistics")
    ch_headers = ["File Name", "Channel", "Minimum", "Maximum", "Mean", "Median", "Std Dev", "Variance"]
    ch_rows = []
    for r in reports:
        if r.get("error"):
            continue
        name = r["general"]["File Name"]
        for ch, s in r.get("channels", {}).get("Per-Channel Statistics", {}).items():
            ch_rows.append([
                name, ch,
                s.get("Minimum"), s.get("Maximum"), s.get("Mean"),
                s.get("Median"), s.get("Standard Deviation"), s.get("Variance"),
            ])
    _write_table(ws3, ch_headers, ch_rows)
    ws3.freeze_panes = "A2"

    # ── Sheet: Data Quality ───────────────────────────────────────────────────
    ws4 = wb.create_sheet("Data Quality")
    q_headers = [
        "File Name", "INF Count", "NaN Count", "Blank Cells", "Corrupted Rows",
        "Wrong Column Count Rows", "Negative Values", "Zero Values", "Repeated Values",
    ]
    q_rows = []
    for r in reports:
        if r.get("error"):
            continue
        name = r["general"]["File Name"]
        q = r.get("quality", {})
        q_rows.append([
            name,
            q.get("INF Count", 0), q.get("NaN Count", 0), q.get("Blank Cell Count", 0),
            q.get("Corrupted Rows", 0), q.get("Rows With Wrong Column Count", 0),
            q.get("Negative Values", 0), q.get("Zero Values", 0), q.get("Repeated Value Count", 0),
        ])
    _write_table(ws4, q_headers, q_rows)
    ws4.freeze_panes = "A2"

    # ── Sheet: Observations ───────────────────────────────────────────────────
    ws5 = wb.create_sheet("Observations")
    obs_rows = [[obs] for obs in generate_observations(data)]
    _write_table(ws5, ["Observation"], obs_rows)
    ws5.freeze_panes = "A2"

    wb.save(report_path)
    return report_path


# ==============================================================================
# MAIN
# ==============================================================================

def main() -> None:
    choice = show_menu()

    if choice == "0":
        print()
        print("  Exiting.")
        print()
        sys.exit(0)

    fmt_choice = select_report_format()

    print()
    print(_SEP)
    print()
    print("  Inspection Started...")

    if choice == "1":
        data = inspect_single_file()
    elif choice == "2":
        data = inspect_single_month()
    elif choice == "3":
        data = inspect_entire_year()
    else:
        sys.exit(1)

    if fmt_choice == "1":
        report_path = save_txt_report(data)
    else:
        report_path = save_excel_report(data)

    print()
    print(_SEP)
    print("Inspection Completed Successfully".center(57))
    print(_SEP)
    print()
    print("  Report Saved")
    print(f"  {report_path}")
    print()
    print(_SEP)
    print()


if __name__ == "__main__":
    main()